import os
import re
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
    ListFlowable,
    ListItem,
)


class DocumentGenerator:
    """
    Generates professional resume and cover-letter documents
    from the student's actual profile and M3.2 AI output.
    """

    def __init__(self, output_dir):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    # ---------------------------------------------------------
    # HELPER METHODS
    # ---------------------------------------------------------

    def _safe_filename(self, text):
        """Convert text into a safe filename."""
        text = str(text or "document")
        text = re.sub(r"[^\w\s-]", "", text)
        text = re.sub(r"\s+", "_", text.strip())
        return text[:100] or "document"

    def _list_value(self, value):
        """Safely convert a value into a list."""
        if value is None:
            return []

        if isinstance(value, list):
            return value

        if isinstance(value, str):
            return [value] if value.strip() else []

        return [str(value)]

    def _add_heading_line(self, paragraph):
        """Add a simple horizontal line below a heading."""
        p = paragraph._p

        pPr = p.get_or_add_pPr()

        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")

        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "666666")

        pBdr.append(bottom)
        pPr.append(pBdr)

    def _add_docx_section_heading(self, document, title):
        """Add a professional section heading to a DOCX."""
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)

        run = paragraph.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)

        self._add_heading_line(paragraph)

        return paragraph

    def _add_bullet(self, document, text):
        """Add a bullet point."""
        if not text:
            return

        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.paragraph_format.space_after = Pt(2)

        run = paragraph.add_run(str(text))
        run.font.size = Pt(10)

    # ---------------------------------------------------------
    # RESUME - DOCX
    # ---------------------------------------------------------

    def generate_resume_docx(
        self,
        candidate_profile,
        selected_job,
        customization,
    ):
        """
        Generate a professional customized resume DOCX.
        """

        document = Document()

        # Page setup
        section = document.sections[0]
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

        # Default font
        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10)

        # -----------------------------------------------------
        # HEADER
        # -----------------------------------------------------

        name = candidate_profile.get("name", "Student")

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)

        run = paragraph.add_run(name)
        run.bold = True
        run.font.size = Pt(18)
        run.font.name = "Arial"

        contact_items = []

        email = candidate_profile.get("email")
        phone = candidate_profile.get("phone")

        if email:
            contact_items.append(str(email))

        if phone:
            contact_items.append(str(phone))

        location = candidate_profile.get("location")
        if location:
            contact_items.append(str(location))

        if contact_items:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)

            run = paragraph.add_run(" | ".join(contact_items))
            run.font.size = Pt(9)

        # -----------------------------------------------------
        # PROFESSIONAL SUMMARY
        # -----------------------------------------------------

        summary = customization.get(
            "tailored_resume_summary",
            ""
        )

        if summary:
            self._add_docx_section_heading(
                document,
                "Professional Summary"
            )

            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(5)

            run = paragraph.add_run(str(summary))
            run.font.size = Pt(10)

        # -----------------------------------------------------
        # TECHNICAL SKILLS
        # -----------------------------------------------------

        skills = customization.get("relevant_skills", [])

        if not skills:
            skills = candidate_profile.get("skills", [])

        if skills:
            self._add_docx_section_heading(
                document,
                "Technical Skills"
            )

            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(5)

            skill_text = ", ".join(
                str(skill) for skill in skills
            )

            run = paragraph.add_run(skill_text)
            run.font.size = Pt(10)

        # -----------------------------------------------------
        # EDUCATION
        # -----------------------------------------------------

        education = candidate_profile.get("education", [])

        if education:
            self._add_docx_section_heading(
                document,
                "Education"
            )

            for item in self._list_value(education):
                self._add_bullet(document, item)

        # -----------------------------------------------------
        # PROJECTS
        # -----------------------------------------------------

        projects = customization.get(
            "relevant_projects",
            []
        )

        if not projects:
            projects = candidate_profile.get(
                "projects",
                []
            )

        if projects:
            self._add_docx_section_heading(
                document,
                "Projects"
            )

            improved_bullets = customization.get(
                "improved_bullet_points",
                []
            )

            for project in self._list_value(projects):
                self._add_bullet(document, project)

            for bullet in self._list_value(
                improved_bullets
            ):
                self._add_bullet(document, bullet)

        # -----------------------------------------------------
        # EXPERIENCE
        # -----------------------------------------------------

        experience = customization.get(
            "relevant_experience",
            []
        )

        if not experience:
            experience = candidate_profile.get(
                "experience",
                []
            )

        if experience:
            self._add_docx_section_heading(
                document,
                "Experience"
            )

            for item in self._list_value(experience):
                self._add_bullet(document, item)

        # -----------------------------------------------------
        # CERTIFICATIONS
        # -----------------------------------------------------

        certifications = customization.get(
            "relevant_certifications",
            []
        )

        if not certifications:
            certifications = candidate_profile.get(
                "certifications",
                []
            )

        if certifications:
            self._add_docx_section_heading(
                document,
                "Certifications"
            )

            for item in self._list_value(
                certifications
            ):
                self._add_bullet(document, item)

        # -----------------------------------------------------
        # JOB KEYWORDS
        # -----------------------------------------------------

        keywords = customization.get(
            "recommended_keywords",
            []
        )

        if keywords:
            self._add_docx_section_heading(
                document,
                "Relevant Keywords"
            )

            paragraph = document.add_paragraph()

            keyword_text = ", ".join(
                str(keyword) for keyword in keywords
            )

            run = paragraph.add_run(keyword_text)
            run.font.size = Pt(10)

        # -----------------------------------------------------
        # SAVE
        # -----------------------------------------------------

        job_title = selected_job.get(
            "job_title",
            selected_job.get("jobtitle", "Job")
        )

        filename = (
            f"{self._safe_filename(name)}_"
            f"{self._safe_filename(job_title)}_"
            f"Resume.docx"
        )

        filepath = os.path.join(
            self.output_dir,
            filename
        )

        document.save(filepath)

        return filepath

    # ---------------------------------------------------------
    # COVER LETTER - DOCX
    # ---------------------------------------------------------

    def generate_cover_letter_docx(
        self,
        candidate_profile,
        selected_job,
        customization,
    ):
        """
        Generate a professional customized cover letter DOCX.
        """

        document = Document()

        section = document.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(11)

        name = candidate_profile.get(
            "name",
            "Student"
        )

        email = candidate_profile.get("email", "")
        phone = candidate_profile.get("phone", "")

        job_title = selected_job.get(
            "job_title",
            selected_job.get("jobtitle", "the position")
        )

        company = selected_job.get(
            "company",
            "the company"
        )

        # Header
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(name)
        run.bold = True
        run.font.size = Pt(16)

        if email or phone:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            contact = " | ".join(
                value for value in [email, phone]
                if value
            )

            run = paragraph.add_run(contact)
            run.font.size = Pt(9)

        document.add_paragraph(
            datetime.now().strftime("%d %B %Y")
        )

        document.add_paragraph(
            f"Hiring Manager\n{company}"
        )

        document.add_paragraph(
            f"Subject: Application for {job_title}"
        )

        document.add_paragraph(
            f"Dear Hiring Manager,"
        )

        cover_letter = customization.get(
            "cover_letter",
            ""
        )

        # If the AI output already contains the letter,
        # use it directly instead of inventing content.
        if cover_letter:
            paragraphs = str(
                cover_letter
            ).split("\n")

            for text in paragraphs:
                text = text.strip()

                if text:
                    paragraph = document.add_paragraph(text)
                    paragraph.paragraph_format.space_after = Pt(9)

        document.add_paragraph(
            "Sincerely,"
        )

        document.add_paragraph(name)

        filename = (
            f"{self._safe_filename(name)}_"
            f"{self._safe_filename(job_title)}_"
            f"Cover_Letter.docx"
        )

        filepath = os.path.join(
            self.output_dir,
            filename
        )

        document.save(filepath)

        return filepath

    # ---------------------------------------------------------
    # RESUME - PDF
    # ---------------------------------------------------------

    def generate_resume_pdf(
        self,
        candidate_profile,
        selected_job,
        customization,
    ):
        """
        Generate a professional customized resume PDF.
        """

        name = candidate_profile.get(
            "name",
            "Student"
        )

        job_title = selected_job.get(
            "job_title",
            selected_job.get("jobtitle", "Job")
        )

        filename = (
            f"{self._safe_filename(name)}_"
            f"{self._safe_filename(job_title)}_"
            f"Resume.pdf"
        )

        filepath = os.path.join(
            self.output_dir,
            filename
        )

        doc = SimpleDocTemplate(
            filepath,
            pagesize=A4,
            rightMargin=0.55 * inch,
            leftMargin=0.55 * inch,
            topMargin=0.5 * inch,
            bottomMargin=0.5 * inch,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "ResumeTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            alignment=TA_CENTER,
            spaceAfter=4,
        )

        contact_style = ParagraphStyle(
            "Contact",
            parent=styles["Normal"],
            fontSize=8.5,
            alignment=TA_CENTER,
            spaceAfter=8,
        )

        section_style = ParagraphStyle(
            "Section",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            spaceBefore=7,
            spaceAfter=3,
        )

        body_style = ParagraphStyle(
            "Body",
            parent=styles["BodyText"],
            fontSize=9.2,
            leading=12,
            spaceAfter=3,
        )

        story = []

        story.append(
            Paragraph(
                str(name),
                title_style
            )
        )

        contact = " | ".join(
            str(value)
            for value in [
                candidate_profile.get("email"),
                candidate_profile.get("phone"),
                candidate_profile.get("location"),
            ]
            if value
        )

        if contact:
            story.append(
                Paragraph(
                    contact,
                    contact_style
                )
            )

        summary = customization.get(
            "tailored_resume_summary",
            ""
        )

        if summary:
            story.append(
                Paragraph(
                    "PROFESSIONAL SUMMARY",
                    section_style
                )
            )

            story.append(
                HRFlowable(
                    width="100%",
                    thickness=0.5
                )
            )

            story.append(
                Paragraph(
                    str(summary),
                    body_style
                )
            )

        skills = customization.get(
            "relevant_skills",
            []
        )

        if not skills:
            skills = candidate_profile.get(
                "skills",
                []
            )

        if skills:
            story.append(
                Paragraph(
                    "TECHNICAL SKILLS",
                    section_style
                )
            )

            story.append(
                Paragraph(
                    ", ".join(
                        str(skill)
                        for skill in skills
                    ),
                    body_style
                )
            )

        self._add_pdf_list_section(
            story,
            "EDUCATION",
            candidate_profile.get("education", []),
            section_style,
            body_style,
        )

        projects = customization.get(
            "relevant_projects",
            []
        )

        if not projects:
            projects = candidate_profile.get(
                "projects",
                []
            )

        self._add_pdf_list_section(
            story,
            "PROJECTS",
            projects,
            section_style,
            body_style,
        )

        self._add_pdf_list_section(
            story,
            "EXPERIENCE",
            customization.get(
                "relevant_experience",
                candidate_profile.get(
                    "experience",
                    []
                )
            ),
            section_style,
            body_style,
        )

        self._add_pdf_list_section(
            story,
            "CERTIFICATIONS",
            customization.get(
                "relevant_certifications",
                candidate_profile.get(
                    "certifications",
                    []
                )
            ),
            section_style,
            body_style,
        )

        self._add_pdf_list_section(
            story,
            "RELEVANT KEYWORDS",
            customization.get(
                "recommended_keywords",
                []
            ),
            section_style,
            body_style,
        )

        doc.build(story)

        return filepath

    def _add_pdf_list_section(
        self,
        story,
        title,
        items,
        section_style,
        body_style,
    ):
        """Add a list-based section to a PDF."""

        items = self._list_value(items)

        if not items:
            return

        story.append(
            Paragraph(
                title,
                section_style
            )
        )

        story.append(
            HRFlowable(
                width="100%",
                thickness=0.5
            )
        )

        list_items = []

        for item in items:
            list_items.append(
                ListItem(
                    Paragraph(
                        str(item),
                        body_style
                    )
                )
            )

        story.append(
            ListFlowable(
                list_items,
                bulletType="bullet",
                leftIndent=15,
            )
        )

    # ---------------------------------------------------------
    # COVER LETTER - PDF
    # ---------------------------------------------------------

    def generate_cover_letter_pdf(
        self,
        candidate_profile,
        selected_job,
        customization,
    ):
        """
        Generate a professional customized cover letter PDF.
        """

        name = candidate_profile.get(
            "name",
            "Student"
        )

        job_title = selected_job.get(
            "job_title",
            selected_job.get("jobtitle", "the position")
        )

        company = selected_job.get(
            "company",
            "the company"
        )

        filename = (
            f"{self._safe_filename(name)}_"
            f"{self._safe_filename(job_title)}_"
            f"Cover_Letter.pdf"
        )

        filepath = os.path.join(
            self.output_dir,
            filename
        )

        doc = SimpleDocTemplate(
            filepath,
            pagesize=A4,
            rightMargin=0.8 * inch,
            leftMargin=0.8 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.75 * inch,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "CoverTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=16,
            alignment=TA_CENTER,
            spaceAfter=4,
        )

        body_style = ParagraphStyle(
            "CoverBody",
            parent=styles["BodyText"],
            fontSize=10.5,
            leading=15,
            spaceAfter=10,
        )

        story = []

        story.append(
            Paragraph(
                str(name),
                title_style
            )
        )

        contact = " | ".join(
            str(value)
            for value in [
                candidate_profile.get("email"),
                candidate_profile.get("phone"),
                candidate_profile.get("location"),
            ]
            if value
        )

        if contact:
            story.append(
                Paragraph(
                    contact,
                    ParagraphStyle(
                        "CoverContact",
                        parent=body_style,
                        fontSize=8.5,
                        alignment=TA_CENTER,
                    )
                )
            )

        story.append(
            Spacer(1, 18)
        )

        story.append(
            Paragraph(
                datetime.now().strftime("%d %B %Y"),
                body_style
            )
        )

        story.append(
            Paragraph(
                f"Hiring Manager<br/>{company}",
                body_style
            )
        )

        story.append(
            Paragraph(
                f"<b>Subject: Application for {job_title}</b>",
                body_style
            )
        )

        story.append(
            Paragraph(
                "Dear Hiring Manager,",
                body_style
            )
        )

        cover_letter = customization.get(
            "cover_letter",
            ""
        )

        if cover_letter:
            for paragraph_text in str(
                cover_letter
            ).split("\n"):

                paragraph_text = paragraph_text.strip()

                if paragraph_text:
                    story.append(
                        Paragraph(
                            paragraph_text,
                            body_style
                        )
                    )

        story.append(
            Paragraph(
                "Sincerely,<br/>" + str(name),
                body_style
            )
        )

        doc.build(story)

        return filepath